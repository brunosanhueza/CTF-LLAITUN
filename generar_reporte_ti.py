import os
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    os.system("pip install python-docx")
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()

    # Título Principal
    title = doc.add_heading('Informe General de Endpoints, Desafíos y Credenciales', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Generado automáticamente para la auditoría del proyecto Aguas del Valle CTF.')

    # 1. Endpoints de la Aplicación Web
    doc.add_heading('1. Endpoints de la Aplicación Web (SCADA HMI)', level=1)
    
    endpoints = [
        ("/", "GET", "Redirección a la vista de inicio de sesión (/login)."),
        ("/login", "GET, POST", "Autenticación de operadores y administradores mediante base de datos."),
        ("/logout", "GET", "Destruye la sesión criptográfica actual del usuario."),
        ("/portal", "GET", "Panel principal del operador. Muestra bitácoras, datos del nodo e información Modbus."),
        ("/portal/upload", "POST", "Sube archivos de mantenimiento y bitácoras. (Contiene vulnerabilidad)."),
        ("/admin", "GET", "Panel exclusivo de Ingeniería SCADA. Revela IPs de la red OT y configuración de actuadores."),
        ("/dashboard", "GET", "Kiosco público (HMI). Representación gráfica de los 4 estanques de agua usando HTML/CSS/JS."),
        ("/api/telemetria", "GET", "Provee el último estado JSON extraído de los registros Modbus para el cliente web."),
        ("/api/reset", "GET, POST", "Fuerza un restablecimiento simulado en caso de fallas de estado.")
    ]

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Endpoint'
    hdr_cells[1].text = 'Método'
    hdr_cells[2].text = 'Descripción'

    for ep, method, desc in endpoints:
        row = table.add_row().cells
        row[0].text = ep
        row[1].text = method
        row[2].text = desc

    doc.add_paragraph()

    # 2. Desafíos Desarrollados (Web TI)
    doc.add_heading('2. Desafíos y Vulnerabilidades Implementadas (Lado TI)', level=1)
    
    doc.add_paragraph('A continuación se detallan las mecánicas de explotación web actuales en el código fuente:')
    
    p1 = doc.add_paragraph()
    p1.add_run('A) Fuga de Código Fuente (.pyc) y Secret Key:\n').bold = True
    p1.add_run('La llave criptográfica de Flask (SECRET_KEY) está quemada (hardcoded) en el archivo app.py. Si el atacante extrae este archivo, puede falsificar (forge) la cookie de sesión para suplantar identidades.')

    p2 = doc.add_paragraph()
    p2.add_run('B) Carga de Archivos Sin Restricción (Unrestricted File Upload):\n').bold = True
    p2.add_run('El endpoint /portal/upload permite adjuntar extensiones peligrosas (.sh, .py, .dat). Las mismas se guardan en el directorio uploads.')

    p3 = doc.add_paragraph()
    p3.add_run('C) Escalada de Privilegios Local (Vulnerabilidad "Dummy" Actual):\n').bold = True
    p3.add_run('Actualmente el sistema lee el contenido de los archivos subidos. Si encuentra las cadenas mágicas "ELEVATE_ROLE=admin", "ADMIN_OVERRIDE_AUTH" o "import os", asciende automáticamente al usuario al rol "admin" en la base de datos. (Nota: Se recomendó cambiar esto a un vector SSTI o Deserialización Insegura para mayor realismo).')

    doc.add_paragraph()

    # 3. Credenciales Registradas
    doc.add_heading('3. Credenciales en Texto Plano (Hardcoded en Base de Datos)', level=1)

    doc.add_paragraph('Las siguientes credenciales se inyectan en texto plano en la base de datos y/o en el código:')

    doc.add_heading('Credenciales Base de Datos (MySQL)', level=2)
    doc.add_paragraph('- Usuario MySQL: root\n- Contraseña MySQL: zacky5023\n- Puerto: 3306')

    doc.add_heading('Cuentas de Acceso al Sistema SCADA (Tabla "usuarios")', level=2)
    
    table2 = doc.add_table(rows=1, cols=4)
    table2.style = 'Table Grid'
    hdr2 = table2.rows[0].cells
    hdr2[0].text = 'Username'
    hdr2[1].text = 'Password (Texto Plano)'
    hdr2[2].text = 'Rol'
    hdr2[3].text = 'Nombre'

    creds = [
        ('operador', 'operador2026', 'operador', 'Carlos Morales (Técnico Operador)'),
        ('admin_scada', 'Adm1n_Pl4nt4_S3cur3!#', 'admin', 'Ing. Rodrigo Silva (Jefe de Planta)')
    ]

    for user, pwd, rol, name in creds:
        r = table2.add_row().cells
        r[0].text = user
        r[1].text = pwd
        r[2].text = rol
        r[3].text = name

    # Guardar
    report_path = os.path.join(os.path.dirname(__file__), 'REPORTE_TI_CTF.docx')
    doc.save(report_path)
    print(f"Reporte generado exitosamente en {report_path}")

if __name__ == "__main__":
    create_report()
